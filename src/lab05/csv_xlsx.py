import csv
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lab04.io_txt_csv import read_text, ensure_parent_dir

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:

    if Path(csv_path).suffix.lower() != ".csv" or Path(xlsx_path).suffix.lower() != ".xlsx":
        raise ValueError("Неверный тип файла")

    if not Path(csv_path).is_file():
            raise FileNotFoundError(f"Файл не найден: {csv_path}")

    text = read_text(csv_path)
    rows = list(csv.reader(text.splitlines()))

    if not rows or not rows[0] or not any(rows[0]):
        raise ValueError("CSV не содержит заголовка")

    if len(rows) < 2:
        raise ValueError("CSV не содержит данных")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for col in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )

        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(8, max_length + 2)

    # Сохраняем результат
    ensure_parent_dir(xlsx_path)
    wb.save(xlsx_path)


if __name__ == "__main__":
    csv_to_xlsx(
        "data/samples/people.csv",
        "data/out/people.xlsx"
    )
    csv_to_xlsx(
            "data/samples/cities.csv",
            "data/out/cities.xlsx"
        )